"""Regression checks for lab paths, using temporary copies and offline services."""

import csv
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
LABS = sorted(path.name for path in ROOT.glob('LabX_*') if path.is_dir())

# Executed in a fresh interpreter to keep each lab's src package independent.
OFFLINE_RUNNER = r'''
import json, os, runpy, sys
from pathlib import Path
from unittest.mock import MagicMock, patch
lab, mode = sys.argv[1:]
base = Path(lab).resolve()
sys.path.insert(0, str(base.parent if mode == 'module' else base))

smtp = MagicMock()
imap = MagicMock()
imap.search.return_value = ('OK', [b''])
response = MagicMock()
response.text = '<article><header><h1>Example Book</h1></header></article><div class="content-body"><ul><li><a>Chapter One</a></li></ul></div>'
driver = MagicMock()
product = MagicMock()
product.find_element.return_value.text = 'Example product'
product.find_element.return_value.get_attribute.return_value = 'https://example.com/product'
driver.find_elements.side_effect = lambda by, selector: [product] if selector == 'article.product_pod' else []

with patch('smtplib.SMTP', return_value=smtp), patch('imaplib.IMAP4_SSL', return_value=imap), patch('requests.get', return_value=response), patch('time.sleep'):
    if base.name == 'LabX_Advance_1':
        prefix = base.name + '.src' if mode == 'module' else 'src'
        def get_driver(manager):
            manager.driver = driver
            return driver
        with patch(prefix + '.driver_manager.DriverManager.get_driver', autospec=True, side_effect=get_driver):
            if mode == 'module':
                runpy.run_module(base.name + '.main', run_name='__main__')
            else:
                runpy.run_path(str(base / 'main.py'), run_name='__main__')
        assert driver.get.called and driver.quit.called
    elif mode == 'module':
        runpy.run_module(base.name + '.main', run_name='__main__')
    else:
        runpy.run_path(str(base / 'main.py'), run_name='__main__')
    if base.name == 'LabX_Basic_12':
        assert smtp.send_message.call_count == 4
        assert imap.search.called
    elif base.name == 'LabX_Advance_12':
        assert smtp.send_message.call_count == 2
        assert imap.search.called
'''


class LabPathTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.repo = Path(self.temp.name) / 'repo'
        self.repo.mkdir()
        self.caller = Path(self.temp.name) / 'unrelated'
        self.caller.mkdir()
        self.env = os.environ.copy()
        for key in ('SENDER_EMAIL', 'SENDER_APP_PASSWORD', 'SMTP_SERVER', 'SMTP_PORT',
                    'IMAP_SERVER', 'TEST_RECIPIENT_EMAIL', 'TEST_SMS_PHONE_NUMBER', 'TEST_SMS_CARRIER'):
            self.env.pop(key, None)

    def copy_lab(self, name):
        lab = self.repo / name
        if lab.exists():
            shutil.rmtree(lab)
        shutil.copytree(ROOT / name, lab, ignore=shutil.ignore_patterns(
            '.git', '__pycache__', '.env', 'screenshots', '*.png', 'reports',
            'attachments', 'incoming_attachments', '.~lock.*'))
        # Delete existing outputs so a stale file cannot make a test pass.
        for pattern in ('data/output*', 'documents/output*', 'documents/generated*',
                        'documents/merged*', 'documents/watermarked*', 'documents/extracted*',
                        'data/processed*', 'data/updated*', 'data/sales_data.json'):
            for path in lab.glob(pattern):
                if path.is_file():
                    path.unlink()
        if name.endswith('_12'):
            (lab / '.env').write_text(
                'SENDER_EMAIL=test@example.com\nSENDER_APP_PASSWORD=test_password\n'
                'SMTP_SERVER=smtp.example.com\nSMTP_PORT=587\nIMAP_SERVER=imap.example.com\n'
                'TEST_RECIPIENT_EMAIL=recipient@example.com\n'
                'TEST_SMS_PHONE_NUMBER=1234567890\nTEST_SMS_CARRIER=att\n')
        return lab

    def run_command(self, args, cwd):
        result = subprocess.run([sys.executable, *args], cwd=cwd, env=self.env,
                                capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertNotIn(' - ERROR - ', result.stderr, result.stderr)
        self.assertNotIn(' - CRITICAL - ', result.stderr, result.stderr)
        return result

    def verify_outputs(self, name, lab):
        if name == 'LabX_Advance_1':
            products = json.loads((lab / 'data/scraped_products.json').read_text())
            self.assertEqual(products[0]['name'], 'Example product')
        elif name == 'LabX_Basic_8':
            with (lab / 'data/output_filtered_sales.csv').open() as stream:
                self.assertEqual(list(csv.DictReader(stream))[-1]['OrderID'], 'SUMMARY')
            inventory = json.loads((lab / 'data/output_updated_inventory.json').read_text())
            self.assertTrue(inventory)
        elif name == 'LabX_Basic_10':
            import openpyxl
            book = openpyxl.load_workbook(lab / 'data/output_sales_report.xlsx')
            self.assertGreater(book.active.cell(book.active.max_row, 4).value, 0)
            book.close()
        elif name == 'LabX_Advance_10':
            import openpyxl
            book = openpyxl.load_workbook(lab / 'data/output_report.xlsx')
            self.assertEqual(book['Raw Data']['D2'].value, '=B2*C2')
            self.assertEqual(len(book['Raw Data']._charts), 1)
            book.close()
        elif name == 'LabX_Basic_11':
            from docx import Document
            from PyPDF2 import PdfReader
            self.assertTrue(Document(lab / 'documents/output_report.docx').paragraphs)
            self.assertEqual(len(PdfReader(lab / 'documents/merged_output.pdf').pages),
                             sum(len(PdfReader(lab / 'documents' / p).pages)
                                 for p in ('dummy_doc1.pdf', 'dummy_doc2.pdf')))
        if name.startswith('LabX_Advance_') and name != 'LabX_Advance_1':
            report_dir = lab / ('documents/reports' if name.endswith('_11') else 'reports')
            if name.endswith('_10'):
                return
            reports = list(report_dir.glob('*.json'))
            self.assertTrue(reports, f'No audit report inside {report_dir}')
            audit = json.loads(reports[0].read_text())
            self.assertFalse([entry for entry in audit if entry['status'] in ('ERROR', 'CRITICAL')])
            self.assertTrue(any(entry['status'] == 'SUCCESS' for entry in audit))

    def test_all_lab_entry_points_from_other_directories(self):
        for name in LABS:
            for mode in ('script', 'module'):
                with self.subTest(lab=name, mode=mode):
                    lab = self.copy_lab(name)
                    network_lab = name.endswith(('_1', '_12'))
                    cwd = self.repo if mode == 'module' else self.caller
                    if network_lab:
                        self.run_command(['-c', OFFLINE_RUNNER, str(lab), mode], cwd)
                    elif mode == 'module':
                        self.run_command(['-m', name + '.main'], cwd)
                    else:
                        # A relative script path must still resolve its resources absolutely.
                        entry = os.path.relpath(lab / 'main.py', cwd)
                        self.run_command([entry], cwd)
                    self.verify_outputs(name, lab)
                    self.assertEqual(list(self.caller.iterdir()), [])
                    self.assertFalse((self.repo / 'data').exists())
                    self.assertFalse((self.repo / 'documents').exists())
                    self.assertFalse((self.repo / 'reports').exists())
                    shutil.rmtree(lab)

    def test_excel_generator_and_report_helper(self):
        import openpyxl
        lab = self.copy_lab('LabX_Advance_10')
        (lab / 'data/input_data.xlsx').unlink()
        self.run_command([str(lab / 'create_dummy_data.py')], self.caller)
        book = openpyxl.load_workbook(lab / 'data/input_data.xlsx')
        self.assertEqual(book['SalesData']['A2'].value, 'Laptop')
        self.assertEqual(book['SalesData'].max_row, 10)
        book.close()
        destination = Path(self.temp.name) / 'explicit-output'
        code = '''import sys
sys.path.insert(0, sys.argv[1])
from LabX_Advance_10.src.utils import save_json_report
assert save_json_report({'value': 7}, 'default.json')
assert save_json_report({'value': 9}, 'nested.json', 'data/custom')
assert save_json_report({'value': 11}, 'absolute.json', sys.argv[2])
'''
        self.run_command(['-c', code, str(self.repo), str(destination)], self.caller)
        self.assertEqual(json.loads((lab / 'data/default.json').read_text()), {'value': 7})
        self.assertTrue((lab / 'data/custom/nested.json').exists())
        self.assertTrue((destination / 'absolute.json').exists())
        self.assertEqual(list(self.caller.iterdir()), [])

    def test_env_is_local_and_external_values_take_precedence(self):
        for name in ('LabX_Basic_12', 'LabX_Advance_12'):
            with self.subTest(lab=name):
                lab = self.copy_lab(name)
                (self.repo / '.env').write_text('SENDER_EMAIL=wrong@example.com\n')
                code = '''import os, sys
sys.path.insert(0, sys.argv[1])
os.environ['SMTP_PORT'] = '2525'
module = __import__(sys.argv[2] + '.src.utils', fromlist=['ENV_FILE'])
assert os.getenv('SENDER_EMAIL') == 'test@example.com'
assert os.getenv('SMTP_PORT') == '2525'
assert module.ENV_FILE == sys.argv[3]
'''
                self.run_command(['-c', code, str(self.repo), name, str(lab / '.env')], self.caller)

    def test_imported_data_agent_keeps_reports_inside_its_lab(self):
        lab = self.copy_lab('LabX_Advance_8')
        code = '''import sys
sys.path.insert(0, sys.argv[1])
from LabX_Advance_8.src.data_agent import DataAgent
DataAgent({'tasks': []}).run()
'''
        self.run_command(['-c', code, str(self.repo)], self.caller)
        self.assertTrue(list((lab / 'reports').glob('*.json')))
        self.assertEqual(list(self.caller.iterdir()), [])


if __name__ == '__main__':
    unittest.main()
